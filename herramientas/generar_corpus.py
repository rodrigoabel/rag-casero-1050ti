#!/usr/bin/env python3
"""
generar_corpus.py - Genera el corpus de prueba del RAG y el banco de preguntas.

Todo el contenido es FICTICIO: empresas, personas, cifras y contratos inventados para que el modelo
no pueda "saber" las respuestas de antemano. Si responde bien, es porque las recuperó del RAG.

Uso (desde la carpeta del proyecto):
    .venv/bin/python herramientas/generar_corpus.py

Genera:
    corpus/usuario1/...        -> se copia tal cual a <almacenamiento>/usuarios/usuario1/ en el servidor
    corpus/usuario2/...        -> prueba de aislamiento (carpeta de Usuario2)
    corpus/cambios/...     -> versiones modificadas para la fase de actualización (F6)
    preguntas/preguntas.json y preguntas.csv
"""

import csv
import json
import os
import random
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from pptx.util import Pt as PPt

ROOT = Path(__file__).resolve().parents[1]
CORPUS = ROOT / "corpus"
PREGUNTAS = ROOT / "preguntas"

USUARIO1 = CORPUS / "usuario1"
MI_RAG = USUARIO1 / "Mi_RAG"
OFFICE = MI_RAG / "office"
PDF_DIR = MI_RAG / "pdf"
PDF_IMG = MI_RAG / "pdf_imagenes"
VAULT = USUARIO1 / "Obsidian_Vault" / "Usuario1" / "Prueba RAG"
USUARIO2 = CORPUS / "usuario2" / "Mi_RAG"
CAMBIOS = CORPUS / "cambios"

FONT_DIR = Path(os.environ.get("CORPUS_FUENTES", "/usr/share/fonts/TTF"))
FONT_REG = FONT_DIR / "DejaVuSans.ttf"
FONT_BOLD = FONT_DIR / "DejaVuSans-Bold.ttf"
FONT_IT = FONT_DIR / "DejaVuSans-Oblique.ttf"

AVISO = "Documento FICTICIO generado para pruebas del RAG doméstico (octubre de 2026). No corresponde a hechos reales."


# =====================================================================================
# D1 - DOCX (office/): Especificación técnica Plataforma Cóndor (informática)
# =====================================================================================
def docx_condor(path: Path, version="2.3", fecha="12 de agosto de 2026", p95="350", jwt_min="45"):
    d = Document()
    d.core_properties.author = "Equipo Cóndor"
    d.core_properties.title = f"Especificación Técnica – Plataforma Cóndor v{version}"
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "Confidencial – Uso interno – Cooperativa Agrícola Valle del Elqui"
    sec.footer.paragraphs[0].text = AVISO

    d.add_heading(f"Especificación Técnica – Plataforma Cóndor v{version}", 0)
    d.add_paragraph(f"Fecha: {fecha}. Cliente: Cooperativa Agrícola Valle del Elqui. "
                    "Arquitecta responsable: Marcela Ibáñez. Estado del documento: aprobado.")

    d.add_heading("1. Propósito", 1)
    d.add_paragraph(
        "La Plataforma Cóndor permite trazar los lotes de uva pisquera desde la recepción en planta hasta el "
        "despacho, emitir certificados de origen y entregar información a los organismos fiscalizadores. "
        "Este documento describe la arquitectura, los requisitos no funcionales y las reglas de seguridad "
        "que deben cumplir todos los componentes del sistema.")

    d.add_heading("2. Arquitectura", 1)
    d.add_paragraph(
        "Todas las solicitudes externas ingresan por un API Gateway Kong que escucha en el puerto 8443 con TLS. "
        "Detrás del gateway operan tres servicios: «lotes», «certificados» y «despachos». La comunicación "
        "asíncrona entre servicios usa RabbitMQ; la cola principal se llama condor.despachos.v1. "
        "La base de datos es PostgreSQL 16, alojada en el servidor «tórtola» del centro de datos de La Serena. "
        "Las consultas frecuentes se guardan en una caché Redis con un tiempo de vida (TTL) de 900 segundos.")

    d.add_heading("3. Requisitos no funcionales", 1)
    for t in [
        "Disponibilidad mensual mínima: 99,7 % (excluye ventanas de mantención programadas).",
        f"Latencia p95 para consultas de lotes: {p95} ms medidos en el gateway.",
        "Objetivo de punto de recuperación (RPO): 15 minutos.",
        "Objetivo de tiempo de recuperación (RTO): 2 horas.",
        f"El token de acceso JWT expira a los {jwt_min} minutos; el token de refresco dura 7 días.",
    ]:
        d.add_paragraph(t, style="List Bullet")

    d.add_heading("4. Ambientes", 1)
    t = d.add_table(rows=1, cols=4, style="Table Grid")
    for c, h in zip(t.rows[0].cells, ["Ambiente", "URL", "Réplicas", "Base de datos"]):
        c.text = h
    for row in [("Desarrollo", "dev.condor.local", "1", "condor_dev"),
                ("QA", "qa.condor.local", "2", "condor_qa"),
                ("Producción", "api.condor.example", "4", "condor_prod")]:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = v

    d.add_heading("5. Endpoints principales", 1)
    t = d.add_table(rows=1, cols=4, style="Table Grid")
    for c, h in zip(t.rows[0].cells, ["Método", "Ruta", "Descripción", "Límite (solicitudes/min)"]):
        c.text = h
    for row in [("GET", "/v2/lotes/{id}", "Consulta un lote", "600"),
                ("POST", "/v2/despachos", "Registra un despacho", "120"),
                ("GET", "/v2/certificados/{folio}", "Obtiene un certificado de origen", "300"),
                ("DELETE", "/v2/lotes/{id}", "Elimina un lote (solo rol supervisor)", "30")]:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = v

    d.add_heading("6. Seguridad", 1)
    d.add_paragraph(
        "Existen tres roles: operador, supervisor y auditor. El rol auditor es de solo lectura. "
        "Solo el rol supervisor puede eliminar lotes. Todas las operaciones quedan en un registro de "
        "auditoría que se retiene durante 400 días y no puede modificarse.")

    d.add_heading("7. Riesgos", 1)
    d.add_paragraph(
        "R-04: dependencia del proveedor de firma electrónica FirmaSur. Si FirmaSur no está disponible, no se "
        "pueden emitir certificados de origen. Mitigación: cola de certificados pendientes y contrato con "
        "un segundo proveedor antes de diciembre de 2026.")
    d.add_paragraph(
        "R-07: crecimiento del volumen de despachos en temporada alta (febrero a abril). Mitigación: "
        "escalamiento horizontal del servicio de despachos y revisión de índices de la base de datos.")

    d.add_heading("8. Historial de versiones", 1)
    t = d.add_table(rows=1, cols=3, style="Table Grid")
    for c, h in zip(t.rows[0].cells, ["Versión", "Fecha", "Cambio principal"]):
        c.text = h
    hist = [("2.1", "3 de mayo de 2026", "Latencia p95 objetivo de 500 ms"),
            ("2.2", "20 de junio de 2026", "Latencia p95 objetivo de 420 ms; se agrega Redis"),
            ("2.3", "12 de agosto de 2026", "Latencia p95 objetivo de 350 ms; token JWT de 45 minutos")]
    if version == "2.4":
        hist.append(("2.4", fecha, f"Latencia p95 objetivo de {p95} ms; token JWT de {jwt_min} minutos"))
    for row in hist:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
    d.save(path)


# =====================================================================================
# D2 - PPTX (office/): Retrospectiva Sprint 42 (informática) - con notas del orador y gráfico
# =====================================================================================
def pptx_sprint(path: Path):
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = "Retrospectiva Sprint 42 – Equipo Cóndor"
    s.placeholders[1].text = "Período: 1 al 14 de septiembre de 2026\n" + AVISO

    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Métricas del sprint"
    tf = s.placeholders[1].text_frame
    tf.text = "Velocidad: 38 puntos (se comprometieron 45)"
    for t in ["Bugs cerrados: 17", "Cobertura de pruebas: 81 %", "Despliegues a producción: 3"]:
        tf.add_paragraph().text = t
    s.notes_slide.notes_text_frame.text = (
        "La velocidad bajó porque un desarrollador estuvo con licencia médica durante 6 días hábiles.")

    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Estado de las historias"
    rows = [("ID", "Historia", "Estado", "Responsable"),
            ("CON-311", "Exportar certificados en PDF", "Terminada", "Javiera Soto"),
            ("CON-318", "Reintentos en la cola de despachos", "En curso", "Tomás Rivas"),
            ("CON-322", "Panel de auditoría", "Bloqueada", "Marcela Ibáñez"),
            ("CON-325", "Migración a PostgreSQL 16", "Terminada", "Tomás Rivas")]
    tbl = s.shapes.add_table(len(rows), 4, Inches(0.4), Inches(1.6), Inches(9.2), Inches(3)).table
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            tbl.cell(i, j).text = v
            tbl.cell(i, j).text_frame.paragraphs[0].font.size = PPt(14)
    s.notes_slide.notes_text_frame.text = (
        "La historia CON-322 está bloqueada porque el proveedor FirmaSur no entrega el certificado de pruebas. "
        "Se escaló a la gerencia el 10 de septiembre de 2026.")

    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "¿Qué salió bien? ¿Qué mejorar?"
    tf = s.placeholders[1].text_frame
    tf.text = "Bien: la migración a PostgreSQL 16 terminó sin caídas"
    for t in ["Bien: buena coordinación con QA",
              "Mejorar: historias sin criterios de aceptación claros",
              "Mejorar: demasiadas reuniones los lunes"]:
        tf.add_paragraph().text = t

    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Bugs por módulo"
    cd = CategoryChartData()
    cd.categories = ["lotes", "despachos", "certificados"]
    cd.add_series("Bugs Sprint 42", (6, 8, 3))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(4.5), cd)

    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Acuerdos del equipo"
    tf = s.placeholders[1].text_frame
    tf.text = "Code freeze los jueves a las 18:00"
    for t in ["Daily a las 9:15 (máximo 15 minutos)",
              "Toda historia nueva debe tener criterios de aceptación antes de entrar al sprint"]:
        tf.add_paragraph().text = t
    p.save(path)


# =====================================================================================
# Utilidades PDF (fpdf2)
# =====================================================================================
class PDF(FPDF):
    def __init__(self, cabecera, pie):
        super().__init__(format="A4")
        self.cabecera, self.pie = cabecera, pie
        self.add_font("DejaVu", "", str(FONT_REG))
        self.add_font("DejaVu", "B", str(FONT_BOLD))
        self.add_font("DejaVu", "I", str(FONT_IT))
        self.set_auto_page_break(True, margin=22)
        self.set_margins(22, 20, 22)

    def multi_cell(self, w, h=None, text="", *args, **kwargs):
        # Tras cada bloque, volver al margen izquierdo y a la línea siguiente
        kwargs.setdefault("new_x", "LMARGIN")
        kwargs.setdefault("new_y", "NEXT")
        return super().multi_cell(w, h, text, *args, **kwargs)

    def header(self):
        self.set_font("DejaVu", "", 8)
        self.set_text_color(110)
        self.cell(0, 6, self.cabecera, align="R")
        self.ln(10)
        self.set_text_color(0)

    def footer(self):
        self.set_y(-16)
        self.set_font("DejaVu", "I", 7)
        self.set_text_color(110)
        self.multi_cell(0, 4, f"{self.pie}\nPágina {self.page_no()} de {{nb}}", align="C")
        self.set_text_color(0)

    def h1(self, t):
        self.ln(3)
        self.set_font("DejaVu", "B", 13)
        self.multi_cell(0, 7, t)
        self.ln(1)

    def h2(self, t):
        self.ln(2)
        self.set_font("DejaVu", "B", 11)
        self.multi_cell(0, 6, t)

    def p(self, t):
        self.set_font("DejaVu", "", 10)
        self.multi_cell(0, 5.4, t, align="J")
        self.ln(2)

    def tabla(self, filas, anchos):
        self.set_font("DejaVu", "", 9)
        with self.table(col_widths=anchos, line_height=6, text_align="LEFT") as t:
            for f in filas:
                r = t.row()
                for c in f:
                    r.cell(c)
        self.ln(3)


# =====================================================================================
# D3 - PDF digital (pdf/): Informe jurídico Concesión Ruta Costera Sur (derecho: concesiones)
# Largo, con muchos plazos y montos parecidos (distractores) y datos clave al final.
# =====================================================================================
def pdf_concesion(path: Path):
    pdf = PDF("Estudio Merino & Asociados — Confidencial", AVISO)
    pdf.add_page()
    pdf.set_font("DejaVu", "B", 16)
    pdf.multi_cell(0, 9, "Informe jurídico: Contrato de concesión «Ruta Costera Sur»", align="C")
    pdf.set_font("DejaVu", "", 10)
    pdf.multi_cell(0, 6, "Tramo Puerto Arena – Caleta Brava (142 km)\nSantiago, 5 de septiembre de 2026", align="C")
    pdf.ln(4)

    pdf.h1("I. Objeto del informe")
    pdf.p("Se nos ha solicitado analizar la situación contractual de la Sociedad Concesionaria Costera Sur S.A. "
          "(en adelante, «la Concesionaria») respecto del contrato de concesión de la obra pública fiscal "
          "denominada «Ruta Costera Sur», con especial énfasis en los plazos, garantías, régimen tarifario, "
          "niveles de servicio, mecanismos de solución de controversias y contingencias vigentes. El análisis "
          "se basa en las bases de licitación, la oferta adjudicada, el decreto de adjudicación y la "
          "correspondencia sostenida con la Inspección Fiscal hasta el 31 de agosto de 2026.")
    pdf.p("Este informe no constituye una opinión sobre aspectos técnicos de ingeniería ni financieros, los "
          "cuales se citan únicamente en la medida en que tienen efectos jurídicos para la Concesionaria.")

    pdf.h1("II. Antecedentes de la adjudicación")
    pdf.p("La concesión fue adjudicada mediante el Decreto Supremo MOP N° 4.417 de 2019 al consorcio que "
          "posteriormente constituyó la Concesionaria. El factor de licitación fue el menor Ingreso Total de "
          "la Concesión (ITC), y la oferta adjudicada fue de UF 21.300.000. Conforme a las bases, el ITC "
          "constituye el monto que, una vez recaudado en valor presente, pone término anticipado a la concesión.")
    pdf.p("La sociedad concesionaria se constituyó dentro del plazo de 60 días contado desde la publicación del "
          "decreto de adjudicación, con un capital suscrito de UF 3.100.000, de los cuales a la fecha se "
          "encuentra pagado el 100 %.")
    pdf.p("Durante la etapa de construcción se suscribió el Convenio Complementario N° 1 (2022), que incorporó "
          "dos enlaces a desnivel y una pasarela peatonal, compensados mediante un aumento del ITC. El "
          "Convenio Complementario N° 2, relativo a la estabilización de taludes, se encuentra en tramitación "
          "(ver sección VII).")

    pdf.h1("III. Plazos del contrato")
    pdf.p("El contrato distingue varios plazos que no deben confundirse entre sí:")
    pdf.tabla([("Plazo", "Duración", "Cómputo"),
               ("Constitución de la sociedad", "60 días", "Desde la publicación del decreto"),
               ("Construcción", "48 meses", "Desde la entrega de terrenos"),
               ("Concesión (máximo)", "32 años", "Desde la puesta en servicio provisoria (PSP)"),
               ("Vigencia garantía de explotación", "Plazo de la concesión + 12 meses", "Desde la PSP"),
               ("Respuesta del Panel Técnico", "30 días", "Desde la presentación de la discrepancia")],
              (48, 45, 70))
    pdf.p("La puesta en servicio provisoria (PSP) de la totalidad de la obra fue autorizada el 15 de marzo de "
          "2024. En consecuencia, el plazo máximo de la concesión vence en marzo de 2056, sin perjuicio de su "
          "término anticipado si se alcanza el ITC antes de esa fecha.")
    pdf.p("Hacemos presente que existieron atrasos menores en la construcción (aproximadamente 4 meses) que "
          "fueron justificados por la Inspección Fiscal debido a lluvias extraordinarias, por lo que no "
          "procedieron multas por atraso.")

    pdf.h1("IV. Garantías")
    pdf.p("Durante la construcción, la Concesionaria entregó una garantía de construcción por un total de "
          "UF 250.000, constituida por 10 boletas de garantía bancarias de UF 25.000 cada una, que fueron "
          "devueltas tras la PSP. Para la etapa de explotación se exige una garantía de explotación de "
          "UF 185.000, cuya vigencia debe cubrir el plazo de la concesión más 12 meses. La renovación de las "
          "boletas debe efectuarse con al menos 90 días de anticipación a su vencimiento; su falta de "
          "renovación faculta al MOP para hacerlas efectivas.")
    pdf.p("Adicionalmente, la Concesionaria mantiene pólizas de seguro de responsabilidad civil por "
          "UF 15.000 por evento y de catástrofe por el valor de reposición de la obra.")

    pdf.h1("V. Régimen tarifario")
    pdf.p("Las tarifas base vigentes para 2026, expresadas en pesos chilenos por pasada en cada plaza de "
          "peaje troncal, son las siguientes:")
    pdf.tabla([("Categoría", "Tipo de vehículo", "Tarifa base (CLP)", "Factor punta"),
               ("1", "Autos y camionetas", "2.300", "1,5"),
               ("2", "Buses", "4.100", "1,0"),
               ("3", "Camiones de 2 ejes", "4.600", "1,0"),
               ("4", "Camiones de más de 2 ejes", "7.900", "1,0"),
               ("5", "Motocicletas", "690", "1,0")],
              (24, 55, 40, 30))
    pdf.p("Las tarifas se reajustan el 1 de enero de cada año según la variación del IPC. Además, el contrato "
          "contempla un premio por seguridad vial de hasta un 5 % de la tarifa si el índice de accidentabilidad "
          "del año anterior es inferior a 0,32. El factor punta solo se aplica en los horarios de punta "
          "definidos por el MOP para los fines de semana largos.")
    pdf.p("El sistema de cobro actual es de plazas de peaje manuales y con telepeaje (TAG). Su eventual "
          "conversión a flujo libre (free-flow) requiere un convenio complementario, materia que se aborda en "
          "las recomendaciones.")

    pdf.h1("VI. Niveles de servicio y régimen de multas")
    pdf.p("La Concesionaria debe mantener un índice de serviciabilidad del pavimento igual o superior a 0,85 "
          "en todos los tramos. El incumplimiento de dicho índice se sanciona con una multa de 40 UTM por cada "
          "día de incumplimiento y por cada tramo afectado.")
    pdf.p("Otras multas relevantes son: no reparar baches en 72 horas (10 UTM por bache y día de atraso) y no "
          "entregar los informes mensuales de explotación (5 UTM por día de atraso).")
    pdf.p("Si las multas acumuladas superan las 1.200 UTM en un período de 12 meses, se configura una causal de "
          "incumplimiento grave del contrato, que habilita al MOP para solicitar la declaración de "
          "incumplimiento grave y la eventual licitación de la concesión por el tiempo que reste.")
    pdf.p("A la fecha de este informe, las multas acumuladas en los últimos 12 meses ascienden a 380 UTM, "
          "principalmente por atrasos en reparación de baches en el sector de Caleta Brava.")

    pdf.h1("VII. Contingencias y riesgos")
    pdf.h2("1. Riesgo geotécnico en el km 47")
    pdf.p("En julio de 2025 se produjo un deslizamiento de talud en el km 47 que obligó a cerrar una pista "
          "durante 23 días. La Concesionaria sostiene que se trató de un evento de fuerza mayor no cubierto por "
          "el riesgo geotécnico asumido en las bases, y solicitó una compensación de UF 1.200.000 por las "
          "obras de estabilización. Dicha compensación se negocia en el marco del Convenio Complementario N° 2, "
          "actualmente en tramitación.")
    pdf.h2("2. Reclamos de usuarios")
    pdf.p("Se registran 47 reclamos ante el SERNAC durante 2025, ninguno con resolución adversa firme. "
          "Estimamos que el riesgo jurídico asociado es bajo.")
    pdf.h2("3. Solución de controversias")
    pdf.p("Las discrepancias técnicas o económicas pueden someterse al Panel Técnico de Concesiones, que debe "
          "emitir una recomendación fundada en un plazo de 30 días; dicha recomendación no es vinculante. Si la "
          "controversia persiste, la Concesionaria puede recurrir a la Comisión Arbitral, cuyas resoluciones son "
          "obligatorias para las partes. Recomendamos agotar la instancia del Panel Técnico respecto de la "
          "compensación del km 47 antes de diciembre de 2026.")

    pdf.h1("VIII. Conclusiones y recomendaciones")
    pdf.p("1. El contrato se encuentra vigente y sin causales de incumplimiento grave configuradas.")
    pdf.p("2. Se recomienda renovar con 90 días de anticipación las boletas de la garantía de explotación que "
          "vencen en 2027.")
    pdf.p("3. Se recomienda solicitar al MOP la modificación del sistema de cobro a flujo libre (free-flow) "
          "antes del 30 de junio de 2027, mediante un convenio complementario que mantenga el equilibrio "
          "económico del contrato.")
    pdf.p("4. Se recomienda presentar la discrepancia sobre la compensación del km 47 ante el Panel Técnico si "
          "no se logra acuerdo en la tramitación del Convenio Complementario N° 2.")
    pdf.ln(6)
    pdf.set_font("DejaVu", "", 10)
    pdf.multi_cell(0, 5.4, "Sebastián Merino Lagos\nAbogado, socio de Estudio Merino & Asociados\n"
                           "Santiago, 5 de septiembre de 2026")
    pdf.output(str(path))


# =====================================================================================
# D4 - PDF digital (pdf/): Procedimiento operativo de sangrado del horno flash (metalurgia)
# =====================================================================================
def pdf_procedimiento(path: Path):
    pdf = PDF("Fundición Los Arrayanes — PO-FLS-014 rev. 3", AVISO)
    pdf.add_page()
    pdf.set_font("DejaVu", "B", 15)
    pdf.multi_cell(0, 8, "Procedimiento operativo PO-FLS-014\nSangrado de mata y escoria del horno flash", align="C")
    pdf.ln(3)
    pdf.tabla([("Código", "PO-FLS-014"), ("Revisión", "3"), ("Vigente desde", "1 de julio de 2026"),
               ("Aprobado por", "Ing. Patricia Ulloa, Superintendente de Fundición"),
               ("Área", "Fundición Los Arrayanes – Horno flash")], (50, 113))

    pdf.h1("1. Objetivo y alcance")
    pdf.p("Establecer la secuencia segura de sangrado de mata y escoria del horno de fusión flash, minimizando "
          "el riesgo de proyecciones de material fundido, explosiones por contacto con agua y pérdidas de "
          "cobre en la escoria. Aplica a operadores, supervisores y personal de mantención de la nave de fusión.")

    pdf.h1("2. Parámetros de operación")
    pdf.tabla([("Parámetro", "Valor / rango", "Observación"),
               ("Temperatura de la mata", "1.180 – 1.220 °C", "Medir con pirómetro antes de abrir la piquera"),
               ("Temperatura de la escoria", "1.250 – 1.280 °C", "Sobre 1.280 °C informar al supervisor"),
               ("Ley de mata objetivo", "62 % Cu (rango 60 – 64 %)", "Muestra cada sangrado"),
               ("Razón Fe/SiO2 en escoria", "1,4", "Ajustar fundente si se desvía más de 0,1"),
               ("Altura mínima de mata para sangrar", "45 cm", "Medida con varilla de sondaje")],
              (55, 50, 58))

    pdf.h1("3. Seguridad")
    pdf.p("Elementos de protección personal obligatorios: traje aluminizado, careta facial con visor de oro, "
          "guantes de cuero de puño largo, polainas y calzado de seguridad metatarsal.")
    pdf.p("Durante el sangrado se debe mantener una distancia mínima de 3 metros del canal de colada. Como "
          "máximo pueden permanecer 2 personas en la plataforma de sangrado. Está prohibido utilizar "
          "herramientas húmedas o mojadas, porque el contacto del agua con el metal fundido genera vapor de "
          "forma violenta y puede provocar una explosión.")

    pdf.h1("4. Secuencia de sangrado")
    pasos = [
        "Verificar que el canal de colada y las ollas estén secos y precalentados.",
        "Confirmar por radio con la sala de control que la altura de mata es igual o superior a 45 cm.",
        "Despejar la plataforma y señalizar el área con cadena y letrero de «Sangrado en curso».",
        "Verificar el EPP de las 2 personas autorizadas en la plataforma.",
        "Medir la temperatura de la mata con el pirómetro y registrarla en la bitácora.",
        "Posicionar la olla de mata bajo el canal y verificar su nivel de llenado anterior.",
        "Abrir la piquera de mata con lanza de oxígeno, con una presión de oxígeno entre 6 y 8 bar.",
        "Controlar el flujo; si se observan proyecciones, retirarse y avisar al supervisor.",
        "Sangrar la escoria por la piquera de escoria N° 2 hacia ollas de 11 m³.",
        "Tomar la muestra de mata para análisis de ley de cobre.",
        "Cerrar la piquera con tapón de arcilla refractaria usando la máquina taponadora.",
        "Registrar tiempos, temperaturas y toneladas sangradas en la bitácora del turno.",
    ]
    pdf.set_font("DejaVu", "", 10)
    for i, t in enumerate(pasos, 1):
        pdf.multi_cell(0, 5.4, f"{i}. {t}")
    pdf.ln(2)

    pdf.h1("5. Emergencias")
    pdf.p("Ante una perforación de crisol o fuga de material fundido fuera del canal se debe activar la alarma "
          "código ROJO-FLASH y evacuar de inmediato al punto de encuentro P-3, ubicado frente a la sala de "
          "compresores. Nadie debe intentar contener el material fundido.")
    pdf.p("En caso de lesiones por quemadura, aplicar el protocolo de primeros auxilios PA-02 y llamar al "
          "policlínico interno al anexo 5500.")

    pdf.h1("6. Registros")
    pdf.p("Bitácora del turno, registro de temperaturas, registro de muestras de mata y lista de verificación "
          "de EPP. Los registros se conservan durante 5 años.")
    pdf.output(str(path))


# =====================================================================================
# D5 - PPTX (office/): Indicadores Fundición Los Arrayanes Q3 2026 (metalurgia)
# =====================================================================================
def pptx_fundicion(path: Path):
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = "Fundición Los Arrayanes – Indicadores de proceso Q3 2026"
    s.placeholders[1].text = "Gerencia de Operaciones – octubre de 2026\n" + AVISO

    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Resumen del trimestre"
    tf = s.placeholders[1].text_frame
    tf.text = "Recuperación de cobre: 97,4 %"
    for t in ["Captura de azufre: 95,8 % (meta: 96,5 %)",
              "Disponibilidad del horno flash: 91,2 %",
              "Producción de ánodos: 41.800 t",
              "Consumo específico de oxígeno: 182 Nm³ por tonelada de concentrado"]:
        tf.add_paragraph().text = t

    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Detalle mensual"
    rows = [("Mes", "Ánodos (t)", "Captura de azufre (%)", "Disponibilidad horno flash (%)"),
            ("Julio", "13.700", "96,1", "92,0"),
            ("Agosto", "14.300", "95,2", "89,8"),
            ("Septiembre", "13.800", "96,0", "91,8")]
    tbl = s.shapes.add_table(len(rows), 4, Inches(0.4), Inches(1.6), Inches(9.2), Inches(2.5)).table
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            tbl.cell(i, j).text = v
            tbl.cell(i, j).text_frame.paragraphs[0].font.size = PPt(14)
    s.notes_slide.notes_text_frame.text = (
        "La caída de agosto en la captura de azufre se debió a la detención de la planta de ácido durante "
        "52 horas por una falla en el soplador principal.")

    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Seguridad"
    tf = s.placeholders[1].text_frame
    tf.text = "2 incidentes con tiempo perdido"
    for t in ["0 fatalidades", "Índice de frecuencia: 1,9"]:
        tf.add_paragraph().text = t

    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Próximos pasos"
    tf = s.placeholders[1].text_frame
    tf.text = "Cambio de refractario del convertidor Peirce-Smith N° 3 en noviembre de 2026 (18 días de detención)"
    for t in ["Auditoría del sistema de captura de gases en diciembre",
              "Meta Q4: captura de azufre de 96,5 %"]:
        tf.add_paragraph().text = t
    p.save(path)


# =====================================================================================
# D6 - DOCX (office/): Informe de clima laboral 2026 (psicología organizacional)
# =====================================================================================
def docx_clima(path: Path):
    d = Document()
    d.core_properties.author = "Ps. Daniela Fuentes"
    d.core_properties.title = "Informe de clima laboral 2026 – Fundición Los Arrayanes"
    d.sections[0].footer.paragraphs[0].text = AVISO
    d.add_heading("Informe de clima laboral 2026 – Fundición Los Arrayanes", 0)
    d.add_paragraph("Elaborado por: Ps. Daniela Fuentes, psicóloga organizacional. Fecha: 22 de septiembre de 2026.")

    d.add_heading("1. Metodología", 1)
    d.add_paragraph(
        "Se aplicó una encuesta en línea entre el 18 y el 29 de agosto de 2026. Respondieron 214 de los 301 "
        "trabajadores, lo que equivale a una tasa de respuesta del 71 %. Se usó una escala Likert de 1 a 5, "
        "donde 5 es la evaluación más favorable. Para resguardar el anonimato, no se reportan resultados de "
        "grupos de menos de 5 personas.")

    d.add_heading("2. Resultados por dimensión", 1)
    t = d.add_table(rows=1, cols=3, style="Table Grid")
    for c, h in zip(t.rows[0].cells, ["Dimensión", "Puntaje promedio (1 a 5)", "Comparación con 2025"]):
        c.text = h
    for row in [("Liderazgo", "3,8", "+0,2"), ("Comunicación interna", "3,1", "-0,1"),
                ("Reconocimiento", "2,7", "-0,3"), ("Seguridad psicológica", "3,4", "+0,1"),
                ("Carga de trabajo", "2,9", "0,0"), ("Compromiso", "4,0", "+0,1")]:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
    d.add_paragraph(
        "La dimensión peor evaluada fue Reconocimiento (2,7) y la mejor evaluada fue Compromiso (4,0). "
        "El eNPS (employee Net Promoter Score) fue de +12.")

    d.add_heading("3. Resultados por área", 1)
    d.add_paragraph(
        "El área de convertidores obtuvo el puntaje más bajo en Reconocimiento (2,6). En esa área predomina el "
        "sistema de turnos 4x4 y los trabajadores señalan que los logros del turno de noche pasan "
        "desapercibidos. El área de refinería obtuvo el mejor resultado en Liderazgo (4,1).")

    d.add_heading("4. Citas representativas (anónimas)", 1)
    d.add_paragraph("«Acá uno se entera de los cambios por el pasillo.» (Comunicación interna)", style="List Bullet")
    d.add_paragraph("«Si el turno sale bien nadie dice nada; si sale mal, se sabe al tiro.» (Reconocimiento)",
                    style="List Bullet")

    d.add_heading("5. Recomendaciones", 1)
    for t in ["Implementar el programa de reconocimiento «Cobre que Suma», con reconocimientos trimestrales "
              "por área y por turno.",
              "Formar a todos los supervisores en entrega de feedback (curso de 16 horas).",
              "Aplicar una encuesta de pulso cada 3 meses para monitorear Reconocimiento y Comunicación interna.",
              "Publicar los cambios operativos en un tablero visible en la sala de cambio de turno."]:
        d.add_paragraph(t, style="List Number")
    d.save(path)


# =====================================================================================
# D7 - PDF escaneado sin texto (pdf_imagenes/): Acta de taller de liderazgo (psicología) -> OCR
# =====================================================================================
def pdf_escaneado_acta(path: Path):
    random.seed(42)
    paginas_txt = [
        ["ACTA DE TALLER DE LIDERAZGO PARA SUPERVISORES",
         "Fundición Los Arrayanes – 3 de septiembre de 2026",
         "",
         "Facilitadora: Ps. Daniela Fuentes",
         "Asistentes: 14 supervisores de las áreas de fusión,",
         "convertidores, refinería y mantención.",
         "",
         "Temas tratados:",
         "1. Resultados preliminares del estudio de clima 2026.",
         "2. Técnicas de feedback y reconocimiento en terreno.",
         "3. Comunicación de cambios operativos al turno.",
         "",
         "Acuerdos:",
         "a) Implementar reuniones 1:1 quincenales de 20 minutos",
         "   entre cada supervisor y su equipo directo.",
         "b) Instalar un tablero de reconocimiento en la sala",
         "   de cambio de turno.",
         "c) Cada supervisor entregará al menos 3 reconocimientos",
         "   por semana a integrantes de su equipo."],
        ["Próxima sesión: 1 de octubre de 2026, sala Ámbar.",
         "",
         "Observaciones de la facilitadora:",
         "El grupo mostró alta disposición. Se sugiere incluir",
         "a los jefes de turno de noche en la próxima sesión,",
         "ya que no pudieron asistir por el sistema de turnos 4x4.",
         "",
         "Firma facilitadora: Daniela Fuentes",
         "",
         AVISO[:52],
         AVISO[52:]],
    ]
    W, H = 1654, 2339  # A4 a 200 ppp
    font = ImageFont.truetype(str(FONT_REG), 34)
    font_b = ImageFont.truetype(str(FONT_BOLD), 36)
    imgs = []
    for lineas in paginas_txt:
        img = Image.new("L", (W, H), 246)
        dr = ImageDraw.Draw(img)
        y = 170
        for i, ln in enumerate(lineas):
            dr.text((150, y), ln, fill=25, font=font_b if (i == 0 and ln.isupper()) else font)
            y += 62
        for _ in range(2500):  # motas de "escáner"
            x, yy = random.randint(0, W - 1), random.randint(0, H - 1)
            img.putpixel((x, yy), random.randint(150, 215))
        img = img.rotate(0.7, resample=Image.BICUBIC, fillcolor=246).filter(ImageFilter.GaussianBlur(0.6))
        imgs.append(img.convert("RGB"))
    imgs[0].save(path, "PDF", resolution=200, save_all=True, append_images=imgs[1:])


# =====================================================================================
# D8 - DOCX en la RAÍZ de Mi_RAG (lo procesa AnythingLLM directamente, sin MarkItDown)
# =====================================================================================
def docx_minuta(path: Path):
    d = Document()
    d.core_properties.title = "Minuta de reunión – Concesión Ruta Costera Sur"
    d.sections[0].footer.paragraphs[0].text = AVISO
    d.add_heading("Minuta de reunión de seguimiento – Concesión Ruta Costera Sur", 0)
    d.add_paragraph("Fecha: 20 de septiembre de 2026. Lugar: oficinas de la Concesionaria, Puerto Arena.")
    d.add_paragraph("Asistentes: Sebastián Merino (abogado), Carolina Pizarro (gerenta general de la "
                    "Concesionaria) y Andrés Vidal (inspector fiscal).")
    d.add_heading("Temas tratados", 1)
    d.add_paragraph(
        "Se revisó el estado de tramitación del Convenio Complementario N° 2 sobre la estabilización del km 47. "
        "Se aclaró una duda histórica: el borrador original de las bases de licitación (2018) contemplaba un plazo "
        "máximo de concesión de 30 años; el contrato adjudicado lo fijó finalmente en 32 años.")
    d.add_heading("Acuerdos", 1)
    for t in ["Solicitar al MOP una prórroga del plazo de tramitación del Convenio Complementario N° 2 hasta el "
              "31 de diciembre de 2026.",
              "Encargar a la consultora Vialtec un estudio de factibilidad del sistema free-flow, con un plazo de "
              "45 días.",
              "Preparar la presentación ante el Panel Técnico en caso de no lograr acuerdo sobre la compensación."]:
        d.add_paragraph(t, style="List Number")
    d.add_paragraph("Próxima reunión: 18 de octubre de 2026.")
    d.save(path)


# =====================================================================================
# D9 - Notas de Obsidian (frontmatter, wikilinks, etiquetas, callouts, tareas, tablas, código)
# =====================================================================================
NOTAS = {
    "Proyectos/Proyecto Cóndor.md": """---
tipo: proyecto
estado: en curso
cliente: Cooperativa Agrícola Valle del Elqui
inicio: 2026-03-02
presupuesto_horas: 1840
tags: [proyecto, condor]
---
# Proyecto Cóndor

Plataforma de trazabilidad de exportaciones de uva pisquera. Arquitecta: [[Marcela Ibáñez]].
Decisión de mensajería: [[ADR-007 Colas con RabbitMQ]].

## Estado
- Próximo hito: salida a producción del módulo de certificados el **17 de noviembre de 2026**.
- Riesgo principal: proveedor de firma electrónica FirmaSur. #riesgo

## Tareas
- [x] Migrar a PostgreSQL 16
- [ ] Definir proveedor alternativo de firma electrónica
- [ ] Preparar manual de usuario

> [!note] Recordatorio
> La especificación técnica vigente está en Mi_RAG (Especificación Técnica – Plataforma Cóndor).

> """ + AVISO + "\n",

    "Decisiones/ADR-007 Colas con RabbitMQ.md": """---
tipo: adr
estado: aceptada
fecha: 2026-04-22
decidido_por: Marcela Ibáñez y Tomás Rivas
tags: [adr, condor, arquitectura]
---
# ADR-007: Colas con RabbitMQ

## Contexto
El [[Proyecto Cóndor]] necesita mensajería asíncrona entre los servicios de lotes, certificados y despachos.
El cliente exige que todo el sistema funcione en sus propios servidores (on-premise), sin servicios de nube pública.

## Alternativas evaluadas
| Opción   | A favor                         | En contra                               |
|----------|---------------------------------|-----------------------------------------|
| RabbitMQ | Simple, conocido por el equipo  | Menor rendimiento a gran escala         |
| Kafka    | Altísimo rendimiento            | Operación compleja para un equipo chico |
| AWS SQS  | Administrado                    | Es nube pública: no cumple on-premise   |

## Decisión
Usar **RabbitMQ** por ser simple de operar, conocido por el equipo y compatible con el requisito on-premise.

> [!warning] Cuándo revisar esta decisión
> No usar Kafka hasta superar los 5.000 mensajes por segundo de forma sostenida.

## Consecuencias
Se acepta un techo de rendimiento menor a cambio de simplicidad operativa.

> """ + AVISO + "\n",

    "Personas/Marcela Ibáñez.md": """---
tipo: persona
rol: arquitecta de software
anexo: 4471
tags: [persona, condor]
---
# Marcela Ibáñez

- Rol: arquitecta responsable del [[Proyecto Cóndor]].
- Especialista en PostgreSQL y modelamiento de datos.
- Prefiere las reuniones antes de las 11:00.
- Anexo telefónico: 4471.

> """ + AVISO + "\n",

    "Diario/2026-09-15.md": """# 2026-09-15

## Reuniones
- Reunión con Tomás Rivas: se decidió mover el respaldo nocturno de la base de datos de las 02:00 a las 03:30
  para que no choque con el proceso ETL de la cooperativa.
- Revisión del [[Proyecto Cóndor]] con [[Marcela Ibáñez]].

## Tareas
- [x] Revisar el PR #212 (reintentos de la cola de despachos)
- [ ] Preparar la demo para la cooperativa (viernes)

> """ + AVISO + "\n",

    "Aprendizajes/Postgres - índices parciales.md": """---
tags: [postgres, rendimiento, aprendizaje]
---
# Postgres: índices parciales

En Cóndor, la consulta de despachos pendientes era lenta porque la tabla `despachos` tiene 12 millones de filas
y solo una fracción pequeña está en estado pendiente. Se creó un índice parcial:

```sql
CREATE INDEX idx_despachos_pendientes
    ON despachos (fecha_creacion)
    WHERE estado = 'pendiente';
```

Resultado: la consulta bajó de 1,8 segundos a 40 milisegundos.

> [!tip] Cuándo usarlo
> Cuando las consultas filtran siempre por el mismo valor y ese valor es una fracción pequeña de la tabla.

> """ + AVISO + "\n",

    "Personal/Receta pebre de la abuela.md": """---
tags: [receta, personal]
---
# Receta: pebre de la abuela

## Ingredientes
- 1 atado de cilantro
- 2 cebollas pequeñas
- 3 tomates
- 1 ají cristal
- 2 cucharadas de vinagre de vino blanco
- Aceite de oliva y sal a gusto

## Preparación
Picar todo muy fino, mezclar y dejar reposar 2 horas en el refrigerador antes de servir.
El secreto de la abuela: agregar el vinagre al final.

> """ + AVISO + "\n",

    "Lugares/Sala Ámbar.md": """---
tipo: sala de reuniones
codigo: AMBAR-3
capacidad: 12
piso: 4
videoconferencia: true
tags: [lugar]
---
Ver propiedades de la sala.
""",
}

NOTA_CONDOR_V2_HITO = ("- Próximo hito: salida a producción del módulo de certificados el **1 de diciembre de 2026** "
                       "(se movió desde el 17 de noviembre por el bloqueo de FirmaSur).")


def escribir_notas():
    for rel, txt in NOTAS.items():
        f = VAULT / rel
        f.parent.mkdir(parents=True, exist_ok=True)
        f.write_text(txt, encoding="utf-8")


# =====================================================================================
# Banco de preguntas
# =====================================================================================
P = []


def q(id, ws, tipo, pregunta, esperada, req, fuente, evid=None, prohib=None, rechazo=False, manual=False,
      rapida=False, fase="F2-F3", dif=1, par=None, variante=None, notas=""):
    P.append({
        "id": id, "fase": fase, "workspace": ws, "tipo": tipo, "dificultad": dif, "rapida": rapida,
        "pregunta": pregunta, "respuesta_esperada": esperada, "requeridas": req,
        "prohibidas": prohib or [], "fuente_esperada": fuente, "evidencia": evid or [],
        "espera_rechazo": rechazo, "revision_manual": manual, "par": par, "variante": variante,
        "notas": notas})


F_CON = "Especificacion_Tecnica_Plataforma_Condor"
F_SPR = "Retrospectiva_Sprint_42_Condor"
F_CRS = "Informe_Juridico_Concesion_Ruta_Costera_Sur"
F_MIN = "Minuta_Reunion_Concesion_2026-09-20"
F_PRO = "Procedimiento_Sangrado_Horno_Flash"
F_FUN = "Indicadores_Fundicion_Los_Arrayanes_Q3_2026"
F_CLI = "Informe_Clima_Laboral_2026_Los_Arrayanes"
F_ACT = "Acta_Taller_Liderazgo_Escaneada"

# --- D1 Cóndor (DOCX vía MarkItDown) ---
q("CON-01", "mirag", "factual", "¿En qué puerto escucha el API Gateway de la Plataforma Cóndor?", "8443 (Kong, con TLS)",
  [["8443"]], F_CON, ["8443"], rapida=True)
q("CON-02", "mirag", "factual", "¿Qué base de datos usa la Plataforma Cóndor y en qué servidor está alojada?",
  "PostgreSQL 16 en el servidor «tórtola»", [["postgres"], ["16"], ["tortola"]], F_CON, ["tortola"])
q("CON-03", "mirag", "parafrasis", "¿Cuánto tiempo puede usar un usuario la Plataforma Cóndor antes de tener que volver a autenticarse?",
  "El token de acceso JWT expira a los 45 minutos (el de refresco dura 7 días)", [["45"]], F_CON, ["45 minutos"], dif=2,
  notas="El documento dice 'token de acceso JWT expira'; la pregunta no usa esas palabras.")
q("CON-04", "mirag", "tabla", "¿Cuántas réplicas tiene el ambiente de producción de la Plataforma Cóndor?", "4 réplicas",
  [["4", "cuatro"]], F_CON, ["api.condor.example"], rapida=True,
  notas="MarkItDown escapa los guiones bajos (condor_prod -> condor\\_prod).")
q("CON-05", "mirag", "tabla", "¿Cuál es el límite de solicitudes por minuto del endpoint POST /v2/despachos en Cóndor?",
  "120 solicitudes por minuto", [["120"]], F_CON, ["/v2/despachos"])
q("CON-06", "mirag", "tabla", "En la Plataforma Cóndor, ¿qué rol se necesita para eliminar un lote y cuál es el límite de solicitudes de ese endpoint?",
  "Rol supervisor; 30 solicitudes por minuto", [["supervisor"], ["30"]], F_CON, ["DELETE"], dif=2)
q("CON-07", "mirag", "version", "¿Cuál es la latencia p95 objetivo para consultas de lotes en la versión vigente (2.3) de la especificación de Cóndor?",
  "350 ms (las versiones 2.1 y 2.2 tenían 500 y 420 ms)", [["350"]], F_CON, ["350 ms"], prohib=["500 ms", "420 ms"],
  rapida=True, dif=2, notas="El historial de versiones incluye valores antiguos (distractores).")
q("CON-08", "mirag", "factual", "¿Durante cuántos días se retiene el registro de auditoría de la Plataforma Cóndor?", "400 días",
  [["400"]], F_CON, ["400 días"])
q("CON-09", "mirag", "sintesis", "Resume los requisitos no funcionales de la Plataforma Cóndor: disponibilidad, RPO y RTO.",
  "Disponibilidad 99,7 %, RPO 15 minutos, RTO 2 horas", [["99,7", "99.7"], ["15"], ["2 horas", "dos horas"]], F_CON,
  ["RPO"], dif=2)
q("CON-10", "mirag", "parafrasis", "¿Qué proveedor externo es un riesgo para la Plataforma Cóndor y qué pasa si falla?",
  "FirmaSur (firma electrónica): si no está disponible no se pueden emitir certificados de origen",
  [["firmasur"]], F_CON, ["FirmaSur"], dif=2)
q("CON-11", "mirag", "metadatos", "¿Quién es la arquitecta responsable de la Plataforma Cóndor y para qué cliente es?",
  "Marcela Ibáñez; Cooperativa Agrícola Valle del Elqui", [["marcela"], ["elqui"]], F_CON, ["Marcela"])
q("CON-12", "mirag", "rechazo", "¿Cuál es el presupuesto en pesos chilenos de la Plataforma Cóndor según su especificación técnica?",
  "No aparece en los documentos (no hay presupuesto en pesos)", [], F_CON, rechazo=True, prohib=["$"], rapida=True,
  notas="Debe decir que no está. Inventar una cifra es alucinación.")

# --- D2 Sprint 42 (PPTX vía MarkItDown, con notas del orador y gráfico) ---
q("SPR-01", "mirag", "factual", "¿Cuál fue la velocidad del equipo Cóndor en el Sprint 42 y cuántos puntos se habían comprometido?",
  "38 puntos de 45 comprometidos", [["38"], ["45"]], F_SPR, ["38 puntos"], rapida=True)
q("SPR-02", "mirag", "tabla", "En el Sprint 42, ¿quién es responsable de la historia CON-318 y en qué estado está?",
  "Tomás Rivas, en curso", [["tomas"], ["en curso"]], F_SPR, ["CON-318"])
q("SPR-03", "mirag", "notas_orador", "¿Por qué está bloqueada la historia CON-322 del Sprint 42?",
  "FirmaSur no entrega el certificado de pruebas; se escaló a gerencia el 10 de septiembre de 2026",
  [["firmasur"]], F_SPR, ["FirmaSur"], rapida=True, dif=2, notas="El dato SOLO está en las notas del orador.")
q("SPR-04", "mirag", "notas_orador", "¿Qué explicación se dio para la baja de velocidad en el Sprint 42?",
  "Un desarrollador estuvo con licencia médica 6 días hábiles", [["licencia"]], F_SPR, ["licencia"], dif=2,
  notas="Solo en notas del orador.")
q("SPR-05", "mirag", "grafico", "Según el gráfico de bugs por módulo del Sprint 42, ¿qué módulo tuvo más bugs?",
  "Despachos (8 bugs)", [["despachos"]], F_SPR, ["despachos"], dif=3,
  notas="Prueba si la conversión extrae los datos de un gráfico de PowerPoint.")
q("SPR-06", "mirag", "factual", "¿Qué día y a qué hora es el code freeze acordado en la retrospectiva del Sprint 42?",
  "Los jueves a las 18:00", [["jueves"], ["18"]], F_SPR, ["Code freeze"])
q("SPR-07", "mirag", "factual", "¿Qué cobertura de pruebas tuvo el Sprint 42 del equipo Cóndor?", "81 %", [["81"]], F_SPR, ["81"])
q("SPR-08", "mirag", "calculo", "¿Qué porcentaje de los puntos comprometidos completó el equipo Cóndor en el Sprint 42?",
  "38/45 ≈ 84 %", [["84", "84,4", "84.4", "84,44", "84.44"]], F_SPR, ["38 puntos"], dif=3, manual=True,
  notas="Cálculo: los modelos pequeños suelen fallar. Aceptable 84,4 % (la regla de números no reconoce 84 dentro de 84,44).")

# --- D3 Concesión (PDF digital largo, PyMuPDF4LLM) ---
q("CRS-01", "mirag", "factual", "¿Cuál es el plazo máximo de la concesión Ruta Costera Sur y desde cuándo se cuenta?",
  "32 años desde la puesta en servicio provisoria (PSP)", [["32", "2056"], ["puesta en servicio", "psp"]], F_CRS,
  ["32 años"], rapida=True, dif=2)
q("CRS-02", "mirag", "factual", "¿Cuál es el monto de la garantía de explotación de la concesión Ruta Costera Sur?",
  "UF 185.000", [["185.000", "185000", "185 000", "185 mil"]], F_CRS, ["185.000"], rapida=True)
q("CRS-03", "mirag", "distractor", "¿Cuál fue el monto de la garantía de construcción de la Ruta Costera Sur y en cuántas boletas se dividió?",
  "UF 250.000 en 10 boletas de UF 25.000", [["250.000", "250000", "250 mil"], ["10", "diez"]], F_CRS, ["250.000"], dif=2)
q("CRS-04", "mirag", "tabla", "¿Cuál es la tarifa base para buses en la Ruta Costera Sur?", "$4.100 (categoría 2)",
  [["4.100", "4100"]], F_CRS, ["4.100"], rapida=True)
q("CRS-05", "mirag", "tabla", "En la Ruta Costera Sur, ¿qué categoría de vehículos tiene un factor punta de 1,5?",
  "Categoría 1: autos y camionetas", [["autos", "categoria 1"]], F_CRS, ["Autos y camionetas"], dif=2)
q("CRS-06", "mirag", "factual", "¿Cuánto es la multa diaria por no cumplir el índice de serviciabilidad en la Ruta Costera Sur?",
  "40 UTM por día y por tramo", [["40"], ["utm"]], F_CRS, ["40 UTM"])
q("CRS-07", "mirag", "profundo", "¿Desde qué monto acumulado de multas se configura incumplimiento grave en la Ruta Costera Sur?",
  "Más de 1.200 UTM en 12 meses", [["1.200", "1200"], ["12 meses", "doce meses"]], F_CRS, ["1.200 UTM"], dif=2)
q("CRS-08", "mirag", "factual", "¿En cuántos días debe emitir su recomendación el Panel Técnico de Concesiones?", "30 días",
  [["30"]], F_CRS, ["Panel Técnico"])
q("CRS-09", "mirag", "factual", "¿Qué ocurrió en el km 47 de la Ruta Costera Sur y cuánto se solicita como compensación?",
  "Deslizamiento de talud en julio de 2025; se solicitan UF 1.200.000",
  [["desliz"], ["1.200.000", "1200000", "1,2 millones"]], F_CRS, ["km 47"], dif=2)
q("CRS-10", "mirag", "profundo", "¿Qué recomienda el informe jurídico respecto del sistema de cobro de la Ruta Costera Sur y con qué plazo?",
  "Solicitar al MOP cambiar a flujo libre (free-flow) antes del 30 de junio de 2027",
  [["free", "flujo libre"], ["2027"]], F_CRS, ["30 de junio de 2027"], rapida=True, dif=2,
  notas="Dato al final de un documento largo.")
q("CRS-11", "mirag", "metadatos", "¿Quién firma el informe jurídico de la Ruta Costera Sur y en qué fecha?",
  "Sebastián Merino Lagos, 5 de septiembre de 2026", [["sebastian merino", "merino lagos"], ["septiembre"]], F_CRS,
  ["Merino Lagos"], notas="«Merino & Asociados» es el estudio, no quien firma.")
q("CRS-12", "mirag", "parafrasis", "¿Cuánto tiempo tiene la concesionaria de la Ruta Costera Sur para construir la obra?",
  "48 meses desde la entrega de terrenos", [["48"]], F_CRS, ["48 meses"], dif=2)
q("CRS-13", "mirag", "premisa_falsa", "¿Por qué la concesión Ruta Costera Sur tiene un plazo de 50 años?",
  "Debe corregir la premisa: el plazo máximo es de 32 años, no 50", [["32"]], F_CRS, ["32 años"], manual=True,
  rapida=True, dif=3, notas="Si acepta los 50 años e inventa una razón, es alucinación.")
q("CRS-14", "mirag", "rechazo", "¿Cuál es el tránsito medio diario anual (TMDA) proyectado de la Ruta Costera Sur?",
  "No aparece en los documentos", [], F_CRS, rechazo=True)
q("CRS-15", "mirag", "sintesis", "Enumera las garantías del contrato de la Ruta Costera Sur con sus montos.",
  "Construcción UF 250.000 (devuelta tras la PSP) y explotación UF 185.000",
  [["250.000", "250000"], ["185.000", "185000"]], F_CRS, ["Garantías"], dif=2)
q("CRS-16", "mirag", "idioma", "What is the maximum term of the Ruta Costera Sur concession?", "32 years from provisional commissioning",
  [["32"]], F_CRS, ["32 años"], dif=2, notas="Pregunta en inglés sobre un documento en español.")

# --- D8 Minuta (DOCX en la raíz de Mi_RAG: lo procesa AnythingLLM sin MarkItDown) ---
q("MIN-01", "mirag", "factual", "¿Hasta qué fecha se acordó solicitar la prórroga del Convenio Complementario N° 2 de la Ruta Costera Sur?",
  "Hasta el 31 de diciembre de 2026", [["31"], ["diciembre"]], F_MIN, ["31 de diciembre"], rapida=True)
q("MIN-02", "mirag", "factual", "¿A qué consultora se encargó el estudio de free-flow y con qué plazo?", "Vialtec, 45 días",
  [["vialtec"], ["45"]], F_MIN, ["Vialtec"])
q("MIN-03", "mirag", "version", "¿Cuál era el plazo máximo de la concesión en el borrador original de las bases y cuál quedó en el contrato adjudicado?",
  "30 años en el borrador de 2018; 32 años en el contrato adjudicado", [["30"], ["32"]], F_MIN, ["borrador"], dif=2)
q("MIN-04", "mirag", "factual", "¿Cuándo es la próxima reunión de seguimiento de la concesión Ruta Costera Sur?",
  "18 de octubre de 2026", [["18"], ["octubre"]], F_MIN, ["18 de octubre"])

# --- D4 Procedimiento (PDF digital) ---
q("PRO-01", "mirag", "tabla", "¿Cuál es la ley de mata objetivo en el procedimiento de sangrado del horno flash?",
  "62 % de cobre (rango 60 – 64 %)", [["62"]], F_PRO, ["62 % Cu"], rapida=True)
q("PRO-02", "mirag", "factual", "¿Qué distancia mínima hay que mantener del canal de colada durante el sangrado del horno flash?",
  "3 metros", [["3", "tres"], ["metro"]], F_PRO, ["3 metros"], rapida=True)
q("PRO-03", "mirag", "parafrasis", "¿Por qué no se pueden usar herramientas mojadas al sangrar el horno flash?",
  "Porque el agua en contacto con el metal fundido genera vapor violento y puede provocar una explosión",
  [["explosi"]], F_PRO, ["explosión"], dif=2)
q("PRO-04", "mirag", "tabla", "¿En qué rango debe estar la temperatura de la escoria del horno flash?", "Entre 1.250 y 1.280 °C",
  [["1.250", "1250"], ["1.280", "1280"]], F_PRO, ["1.250"])
q("PRO-05", "mirag", "factual", "¿Qué presión de oxígeno se usa al abrir la piquera de mata con lanza?", "Entre 6 y 8 bar",
  [["6"], ["8"], ["bar"]], F_PRO, ["6 y 8 bar"])
q("PRO-06", "mirag", "factual", "Ante una perforación de crisol en el horno flash, ¿qué alarma se activa y a qué punto de encuentro se evacúa?",
  "Alarma código ROJO-FLASH; punto de encuentro P-3", [["rojo"], ["p-3", "p3", "p 3"]], F_PRO, ["ROJO-FLASH"], rapida=True)
q("PRO-07", "mirag", "metadatos", "¿Quién aprobó el procedimiento PO-FLS-014 y desde cuándo está vigente?",
  "Ing. Patricia Ulloa; vigente desde el 1 de julio de 2026", [["ulloa"], ["julio"]], F_PRO, ["Patricia Ulloa"])
q("PRO-08", "mirag", "factual", "¿Cuál es el código y la revisión del procedimiento de sangrado del horno flash?",
  "PO-FLS-014, revisión 3", [["fls-014", "fls 014", "fls014"], ["3"]], F_PRO, ["PO-FLS-014"])
q("PRO-09", "mirag", "rechazo", "Según el procedimiento de sangrado del horno flash, ¿a qué temperatura se deben moldear los ánodos?",
  "No aparece en el procedimiento", [], F_PRO, rechazo=True, rapida=True)
q("PRO-10", "mirag", "factual", "¿Cuántas personas pueden estar como máximo en la plataforma de sangrado?", "2 personas",
  [["2", "dos"]], F_PRO, ["2 personas"])

# --- D5 Fundición Q3 (PPTX) ---
q("FUN-01", "mirag", "factual", "¿Cuál fue la recuperación de cobre de la Fundición Los Arrayanes en el tercer trimestre de 2026?",
  "97,4 %", [["97,4", "97.4"]], F_FUN, ["97,4"], rapida=True)
q("FUN-02", "mirag", "factual", "¿Se cumplió la meta de captura de azufre de la Fundición Los Arrayanes en Q3 2026? Indica el valor y la meta.",
  "No: 95,8 % frente a una meta de 96,5 %", [["95,8", "95.8"], ["96,5", "96.5"]], F_FUN, ["95,8"], dif=2)
q("FUN-03", "mirag", "tabla", "¿Cuántas toneladas de ánodos produjo la Fundición Los Arrayanes en agosto de 2026?", "14.300 t",
  [["14.300", "14300"]], F_FUN, ["14.300"], rapida=True)
q("FUN-04", "mirag", "notas_orador", "¿Por qué bajó la captura de azufre en agosto de 2026 en la Fundición Los Arrayanes?",
  "La planta de ácido estuvo detenida 52 horas por una falla en el soplador principal", [["52"]], F_FUN, ["52 horas"],
  dif=2, notas="Solo en notas del orador.")
q("FUN-05", "mirag", "factual", "¿Cuál fue el consumo específico de oxígeno de la Fundición Los Arrayanes en Q3 2026?",
  "182 Nm³ por tonelada de concentrado", [["182"]], F_FUN, ["182"])
q("FUN-06", "mirag", "factual", "¿Cuántos incidentes con tiempo perdido hubo en la Fundición Los Arrayanes en Q3 2026 y hubo fatalidades?",
  "2 incidentes con tiempo perdido; 0 fatalidades", [["2", "dos"], ["0", "cero", "ninguna", "sin fatalidades", "no hubo"]],
  F_FUN, ["tiempo perdido"])
q("FUN-07", "mirag", "factual", "¿Qué mantención mayor está planificada en la Fundición Los Arrayanes para noviembre de 2026 y cuántos días durará?",
  "Cambio de refractario del convertidor Peirce-Smith N° 3, 18 días", [["refractario"], ["18"]], F_FUN, ["refractario"])
q("FUN-08", "mirag", "calculo", "¿Cuál fue el promedio mensual de producción de ánodos en Q3 2026 en la Fundición Los Arrayanes?",
  "41.800 / 3 ≈ 13.933 t", [["13.93", "13933", "13 933"]], F_FUN, ["41.800"], dif=3, manual=True)
q("FUN-09", "mirag", "factual", "¿Cuál fue la disponibilidad del horno flash en el tercer trimestre de 2026?", "91,2 %",
  [["91,2", "91.2"]], F_FUN, ["91,2"])

# --- D6 Clima laboral (DOCX) ---
q("CLI-01", "mirag", "factual", "¿Cuántas personas respondieron la encuesta de clima laboral 2026 de la Fundición Los Arrayanes y cuál fue la tasa de respuesta?",
  "214 de 301 trabajadores; 71 %", [["214"], ["71"]], F_CLI, ["214"], rapida=True)
q("CLI-02", "mirag", "parafrasis", "¿Qué porcentaje de los trabajadores de Los Arrayanes contestó el estudio de clima?",
  "71 %", [["71"]], F_CLI, ["71 %"], dif=2)
q("CLI-03", "mirag", "tabla", "¿Cuál fue la dimensión peor evaluada en el estudio de clima 2026 y con qué puntaje?",
  "Reconocimiento, 2,7", [["reconocimiento"], ["2,7", "2.7"]], F_CLI, ["Reconocimiento"], rapida=True)
q("CLI-04", "mirag", "factual", "¿Cuál fue el eNPS del estudio de clima 2026 de Los Arrayanes?", "+12", [["12"]], F_CLI, ["eNPS"])
q("CLI-05", "mirag", "factual", "¿Qué área tuvo el puntaje más bajo en reconocimiento en el estudio de clima?", "Convertidores (2,6)",
  [["convertidor"], ["2,6", "2.6"]], F_CLI, ["convertidores"], dif=2)
q("CLI-06", "mirag", "factual", "¿Qué recomienda el informe de clima laboral para mejorar el reconocimiento?",
  "Programa «Cobre que Suma» con reconocimientos trimestrales (y formar supervisores en feedback, 16 horas)",
  [["cobre que suma"]], F_CLI, ["Cobre que Suma"], dif=2)
q("CLI-07", "mirag", "parafrasis", "¿Por qué el informe de clima no muestra los resultados de algunos equipos?",
  "Para resguardar el anonimato no se reportan grupos de menos de 5 personas", [["5", "cinco"]], F_CLI, ["anonimato"], dif=2)
q("CLI-08", "mirag", "factual", "¿Qué frase textual de un trabajador ilustra los problemas de comunicación interna?",
  "«Acá uno se entera de los cambios por el pasillo.»", [["pasillo"]], F_CLI, ["pasillo"])
q("CLI-09", "mirag", "sintesis", "Resume en 3 puntos los principales hallazgos del informe de clima laboral 2026 de Los Arrayanes.",
  "Debe mencionar al menos: reconocimiento peor evaluado (2,7), compromiso mejor (4,0) y convertidores como área crítica",
  [["reconocimiento"]], F_CLI, ["Reconocimiento"], manual=True, dif=3)
q("CLI-10", "mirag", "rechazo", "¿Cuál es el índice de rotación de personal de la Fundición Los Arrayanes en 2026?",
  "No aparece en los documentos", [], F_CLI, rechazo=True)

# --- D7 Acta escaneada (OCR) ---
q("ACT-01", "mirag", "ocr", "¿Cuántos supervisores asistieron al taller de liderazgo del 3 de septiembre de 2026?", "14",
  [["14", "catorce"]], F_ACT, ["14 supervisores"], rapida=True)
q("ACT-02", "mirag", "ocr", "¿Qué se acordó sobre las reuniones uno a uno en el taller de liderazgo para supervisores?",
  "Reuniones 1:1 quincenales de 20 minutos", [["quincenal"], ["20"]], F_ACT, ["quincenales"], dif=2)
q("ACT-03", "mirag", "ocr", "¿Cuándo y dónde es la próxima sesión del taller de liderazgo para supervisores?",
  "1 de octubre de 2026, sala Ámbar", [["octubre"], ["ambar"]], F_ACT, ["Ámbar"])
q("ACT-04", "mirag", "ocr", "¿Cuántos reconocimientos por semana se comprometió a entregar cada supervisor?", "Al menos 3",
  [["3", "tres"]], F_ACT, ["reconocimientos"])

# --- Cruce de documentos (todos en el workspace de Mi_RAG) ---
q("X-01", "mirag", "cruce", "¿Qué tiene en común el riesgo R-04 de la especificación de Cóndor con la historia bloqueada CON-322?",
  "Ambos dependen del proveedor de firma electrónica FirmaSur", [["firmasur"]], [F_CON, F_SPR], ["FirmaSur"], dif=3)
q("X-02", "mirag", "cruce", "Para la Fundición Los Arrayanes, entrega la captura de azufre de Q3 2026 y la tasa de respuesta de la encuesta de clima.",
  "95,8 % y 71 %", [["95,8", "95.8"], ["71"]], [F_FUN, F_CLI], ["95,8"], dif=3)
q("X-03", "mirag", "cruce", "¿Quién facilitó el taller de liderazgo para supervisores y qué otro documento elaboró esa persona?",
  "Ps. Daniela Fuentes, que también elaboró el informe de clima laboral 2026", [["daniela"], ["clima"]], [F_ACT, F_CLI],
  ["Daniela Fuentes"], dif=3)
q("X-04", "mirag", "cruce", "¿Qué área de la Fundición Los Arrayanes aparece en el estudio de clima y también en la mantención planificada para noviembre?",
  "Convertidores: peor puntaje en reconocimiento (2,6) y cambio de refractario del convertidor Peirce-Smith N° 3",
  [["convertidor"]], [F_CLI, F_FUN], ["convertidor"], manual=True, dif=3)

# --- Obsidian ---
F_OB_CON = "Proyecto-Condor"
q("OBS-01", "obsidian", "frontmatter", "¿Cuántas horas de presupuesto tiene el Proyecto Cóndor?", "1.840 horas",
  [["1840", "1.840"]], F_OB_CON, ["1840"], rapida=True, dif=2, notas="El dato solo está en las propiedades (frontmatter).")
q("OBS-02", "obsidian", "factual", "¿Cuál es el próximo hito del Proyecto Cóndor y en qué fecha?",
  "Salida a producción del módulo de certificados el 17 de noviembre de 2026", [["17"], ["noviembre"]], F_OB_CON,
  ["17 de noviembre"], rapida=True)
q("OBS-03", "obsidian", "sintesis", "¿Por qué se eligió RabbitMQ en lugar de Kafka para el Proyecto Cóndor?",
  "Simple de operar, conocido por el equipo y compatible con el requisito on-premise; Kafka es complejo de operar para un equipo chico",
  [["rabbitmq"]], "ADR-007", ["RabbitMQ"], manual=True, dif=2)
q("OBS-04", "obsidian", "callout", "¿A partir de qué volumen de mensajes se debería reconsiderar Kafka en Cóndor?",
  "5.000 mensajes por segundo sostenidos", [["5.000", "5000", "5 000", "5,000"]], "ADR-007", ["5.000"], dif=2,
  notas="El dato está dentro de un callout > [!warning].")
q("OBS-05", "obsidian", "factual", "¿Cuál es el anexo de Marcela Ibáñez y en qué horario prefiere las reuniones?",
  "Anexo 4471; antes de las 11:00", [["4471"], ["11"]], "Marcela", ["4471"], rapida=True)
q("OBS-06", "obsidian", "factual", "¿Por qué se movió el respaldo nocturno de la base de datos y a qué hora quedó?",
  "Para no chocar con el ETL de la cooperativa; quedó a las 03:30", [["3:30", "03:30"], ["etl"]], "2026-09-15", ["03:30"], dif=2)
q("OBS-07", "obsidian", "tareas", "¿Qué tareas quedaron pendientes en la nota diaria del 15 de septiembre de 2026?",
  "Preparar la demo para la cooperativa (viernes)", [["demo"]], "2026-09-15", ["demo"], dif=2)
q("OBS-08", "obsidian", "factual", "¿Cuánto mejoró la consulta de despachos pendientes con el índice parcial de PostgreSQL?",
  "De 1,8 segundos a 40 milisegundos", [["1,8", "1.8"], ["40"]], "indices-parciales", ["40 milisegundos"])
q("OBS-09", "obsidian", "personal", "¿Cuánto tiempo hay que dejar reposar el pebre de la abuela?", "2 horas en el refrigerador",
  [["2", "dos"], ["hora"]], "pebre", ["reposar"], rapida=True)
q("OBS-10", "obsidian", "frontmatter", "¿Cuál es la capacidad de la sala Ámbar y en qué piso está?", "12 personas, piso 4",
  [["12"], ["4", "cuarto"]], "Sala-Ambar", ["capacidad"], dif=3, notas="La nota solo tiene propiedades.")
q("OBS-11", "obsidian", "rechazo", "¿Qué framework de frontend usa el Proyecto Cóndor?", "No aparece en las notas", [],
  F_OB_CON, rechazo=True, rapida=True)
q("OBS-12", "obsidian", "enlaces", "¿Quiénes decidieron el ADR-007 y qué rol tiene Marcela Ibáñez en el Proyecto Cóndor?",
  "Marcela Ibáñez y Tomás Rivas; Marcela es la arquitecta responsable", [["marcela"], ["tomas"], ["arquitect"]],
  "ADR-007", ["decidido_por"], dif=3, manual=True)

# --- Aislamiento entre usuarios ---
q("AIS-01", "mirag", "aislamiento", "¿Cuál es la palabra clave de aislamiento de Usuario2?", "No debe encontrarla (es un documento de Usuario2)",
  [], "Prueba_Aislamiento_Usuario2", rechazo=True, prohib=["guayacan"], fase="F7", rapida=True)
q("AIS-02", "obsidian", "aislamiento", "¿Cuál es la palabra clave de aislamiento de Usuario2?", "No debe encontrarla",
  [], "Prueba_Aislamiento_Usuario2", rechazo=True, prohib=["guayacan"], fase="F7")
q("AIS-03", "usuario2", "aislamiento", "¿Cuál es la palabra clave de aislamiento de Usuario2?", "GUAYACÁN-77 (control positivo en el espacio de Usuario2)",
  [["guayacan"]], "Prueba_Aislamiento_Usuario2", ["GUAYACÁN"], fase="F7")

# --- Actualización y borrado (solo después de aplicar corpus/cambios, fase F6) ---
q("UPD-01", "mirag", "actualizacion", "¿Cuál es la latencia p95 objetivo para consultas de lotes en la especificación de la Plataforma Cóndor?",
  "300 ms (versión 2.4)", [["300"]], F_CON, ["300 ms"], prohib=["350 ms"], fase="F6",
  notas="Si aparece 350 ms, quedó la versión anterior en el RAG (duplicado).")
q("UPD-02", "mirag", "actualizacion", "¿A los cuántos minutos expira el token de acceso JWT de la Plataforma Cóndor?",
  "30 minutos (versión 2.4)", [["30"]], F_CON, ["30 minutos"], prohib=["45 minutos"], fase="F6")
q("UPD-03", "obsidian", "actualizacion", "¿Cuál es el próximo hito del Proyecto Cóndor y en qué fecha?",
  "1 de diciembre de 2026 (movido desde el 17 de noviembre)", [["diciembre"]], F_OB_CON, ["1 de diciembre"], fase="F6")
q("UPD-04", "obsidian", "borrado", "¿Cuánto tiempo hay que dejar reposar el pebre de la abuela?",
  "Ya no debería encontrarse (la nota fue borrada)", [], "pebre", rechazo=True, prohib=["2 horas"], fase="F6")

# --- Formulación: pregunta ingenua vs bien hecha (fase F4, para aprender a preguntar) ---
pares = [
    ("FORM-01", "¿y el caché?",
     "En la especificación técnica de la Plataforma Cóndor, ¿cuál es el TTL de la caché Redis? Responde con el número y la unidad.",
     "900 segundos (15 minutos)", [["900", "15 minutos"]], F_CON, ["900 segundos"]),
    ("FORM-02", "multa?",
     "Según el informe jurídico de la concesión Ruta Costera Sur, ¿cuál es la multa diaria por incumplir el índice de serviciabilidad? Responde solo con el monto y la unidad.",
     "40 UTM por día", [["40"], ["utm"]], F_CRS, ["40 UTM"]),
    ("FORM-03", "cuéntame del horno",
     "Según el procedimiento PO-FLS-014 de sangrado del horno flash, ¿en qué rango debe estar la temperatura de la mata?",
     "Entre 1.180 y 1.220 °C", [["1.180", "1180"], ["1.220", "1220"]], F_PRO, ["1.180"]),
    ("FORM-04", "¿cómo está el clima?",
     "En el informe de clima laboral 2026 de la Fundición Los Arrayanes, ¿qué puntaje obtuvo la dimensión Comunicación interna?",
     "3,1", [["3,1", "3.1"]], F_CLI, ["Comunicación interna"]),
    ("FORM-05", "dame todo lo de la concesión",
     "Según el informe jurídico de la Ruta Costera Sur, indica en viñetas: plazo máximo de la concesión, plazo de construcción y fecha de puesta en servicio provisoria.",
     "32 años; 48 meses; 15 de marzo de 2024", [["32", "2056"], ["48"], ["2024"]], F_CRS, ["15 de marzo de 2024"]),
    ("FORM-06", "¿qué pasó con la planta de ácido?",
     "En la presentación de indicadores Q3 2026 de la Fundición Los Arrayanes, ¿cuántas horas estuvo detenida la planta de ácido en agosto y por qué?",
     "52 horas, por una falla en el soplador principal", [["52"]], F_FUN, ["52 horas"]),
    ("FORM-07", "¿quién es Tomás?",
     "En la retrospectiva del Sprint 42 del equipo Cóndor, ¿de qué historias es responsable Tomás Rivas y en qué estado está cada una?",
     "CON-318 (en curso) y CON-325 (terminada)", [["318"], ["325"]], F_SPR, ["Tomás Rivas"]),
    ("FORM-08", "¿qué es el ITC?",
     "En el informe jurídico de la Ruta Costera Sur, ¿cuál fue el Ingreso Total de la Concesión (ITC) de la oferta adjudicada?",
     "UF 21.300.000", [["21.300.000", "21300000"]], F_CRS, ["21.300.000"]),
]
for pid, ing, mej, esp, req, fuente, evid in pares:
    q(pid + "a", "mirag", "formulacion", ing, esp, req, fuente, evid, fase="F4", par=pid, variante="ingenua", dif=3)
    q(pid + "b", "mirag", "formulacion", mej, esp, req, fuente, evid, fase="F4", par=pid, variante="mejorada", dif=1)

# --- E8: preguntas ingenuas al estilo de la familia por Telegram (reformulación con otro modelo) ---
# Cada una copia la respuesta esperada, las palabras clave y la evidencia de su versión bien hecha de la F3
# (campo «par»), para comparar ingenua, reformulada y bien hecha. Tienen fase propia: no entran en las corridas
# habituales (--fase E8). Incluyen lo típico de una pregunta real: sin nombre de documento, palabras distintas a
# las del documento, fechas relativas («el trimestre pasado» = Q3 2026, vista desde octubre de 2026), premisas
# falsas, faltas de ortografía y datos que no existen.
ingenuas_e8 = [
    # derecho: concesiones
    ("ING-01", "CRS-01", "cuanto dura la concesion?"),
    ("ING-02", "CRS-02", "monto garantia explotacion"),
    ("ING-03", "CRS-06", "q multa hay si no cumplen el estado del pavimento"),
    ("ING-04", "CRS-07", "cuándo pasa a ser incumplimiento grave?"),
    ("ING-05", "CRS-09", "qué pasó con el derrumbe del km 47"),
    ("ING-06", "CRS-12", "cuánto tienen para terminar las obras?"),
    ("ING-07", "CRS-13", "por qué le dieron 50 años a la concesionaria?"),
    ("ING-08", "MIN-01", "hasta cuándo pidieron la prórroga del convenio?"),
    ("ING-09", "MIN-04", "cuándo nos volvemos a juntar por la concesión?"),
    ("ING-10", "CRS-08", "cuánto se demora el panel en recomendar?"),
    # psicología organizacional
    ("ING-11", "CLI-02", "cuánta gente contestó la encuesta?"),
    ("ING-12", "CLI-03", "qué salió peor en clima?"),
    ("ING-13", "CLI-05", "qué área está peor con el reconocimiento?"),
    ("ING-14", "CLI-06", "qué proponen para que la gente se sienta reconocida?"),
    ("ING-15", "CLI-07", "por qué no aparecen algunos equipos en el informe?"),
    ("ING-16", "ACT-01", "cuántos vinieron al taller?"),
    ("ING-17", "ACT-03", "cuándo es la próxima sesión del taller?"),
    ("ING-18", "CLI-04", "y el enps?"),
    # metalurgia del cobre
    ("ING-19", "PRO-04", "temperatura escoria?"),
    ("ING-20", "PRO-02", "a qué distancia hay que pararse de la canaleta cuando se sangra?"),
    ("ING-21", "PRO-06", "qué hago si se perfora el crisol?"),
    ("ING-22", "PRO-05", "con cuánta presión va el oxígeno para abrir la piquera?"),
    ("ING-23", "FUN-01", "cómo nos fue en recuperación el trimestre pasado?"),
    ("ING-24", "FUN-02", "cumplimos la meta de azufre?"),
    ("ING-25", "FUN-03", "cuántos ánodos se hicieron en agosto?"),
    ("ING-26", "FUN-07", "qué mantención viene en noviembre?"),
    ("ING-27", "FUN-09", "disponibilidad del flash el trimestre pasado"),
    # ingeniería de software
    ("ING-28", "CON-01", "puerto del gateway?"),
    ("ING-29", "CON-08", "cuánto se guardan los logs de auditoría?"),
    ("ING-30", "SPR-07", "cobertura del sprint?"),
    ("ING-31", "OBS-06", "a qué hora quedó el backup de la base?"),
    ("ING-32", "OBS-10", "cuánta gente cabe en la sala ámbar?"),
    # datos que no existen: el reformulador no debe inventar contexto que haga «aparecer» una respuesta
    ("ING-33", "CON-12", "cuánto cuesta cóndor?"),
    ("ING-34", "PRO-09", "a cuántos grados se moldean los ánodos?"),
    ("ING-35", "CLI-10", "cuánta rotación tuvo la fundición este año?"),
]
_por_id = {x["id"]: x for x in P}
for iid, base, texto in ingenuas_e8:
    b = _por_id[base]
    q(iid, b["workspace"], "ingenua", texto, b["respuesta_esperada"], b["requeridas"], b["fuente_esperada"],
      b["evidencia"], prohib=b["prohibidas"], rechazo=b["espera_rechazo"], fase="E8", dif=3, par=base,
      variante="ingenua", notas=f"Versión bien hecha: {base} («{b['pregunta']}»)")

# E8b: comprobación con preguntas que ningún reformulador vio al diseñar sus reglas (9/10/2026, antes de medir la v3).
# Cada una pide lo mismo que verifica su par de la F3, para no repetir los artefactos de ING-11 e ING-32.
ingenuas_e8b = [
    ("ING-36", "CON-02", "qué base de datos usa cóndor y en qué servidor está?"),
    ("ING-37", "CON-04", "cuántas réplicas hay en prod?"),
    ("ING-38", "SPR-06", "cuándo es el code freeze?"),
    ("ING-39", "SPR-03", "por qué está trabada la CON-322?"),
    ("ING-40", "CRS-03", "garantia de construccion cuanto era y en cuantas boletas"),
    ("ING-41", "CRS-10", "qué recomiendan hacer con el cobro de los peajes y para cuándo?"),
    ("ING-42", "MIN-02", "a quién le encargaron lo del free flow y en cuánto tiempo?"),
    ("ING-43", "CRS-11", "quién firmó el informe de la concesión y cuándo?"),
    ("ING-44", "PRO-01", "ley de mata?"),
    ("ING-45", "PRO-10", "cuánta gente puede haber arriba de la plataforma cuando se sangra?"),
    ("ING-46", "FUN-05", "consumo de oxígeno del trimestre pasado?"),
    ("ING-47", "FUN-06", "hubo accidentes en la fundición el trimestre pasado? algún fallecido?"),
    ("ING-48", "ACT-04", "cuántos reconocimientos por semana tiene que dar cada supervisor?"),
    ("ING-49", "OBS-03", "por qué elegimos rabbit y no kafka?"),
    ("ING-50", "CRS-14", "cuántos vehículos al día se proyectan para la ruta costera?"),
    ("ING-51", "MIN-03", "por qué al final la concesión quedó en 30 años?"),
]
for iid, base, texto in ingenuas_e8b:
    b = _por_id[base]
    q(iid, b["workspace"], "ingenua", texto, b["respuesta_esperada"], b["requeridas"], b["fuente_esperada"],
      b["evidencia"], prohib=b["prohibidas"], rechazo=b["espera_rechazo"], fase="E8b", dif=3, par=base,
      variante="ingenua", notas=f"Comprobación E8b. Versión bien hecha: {base} («{b['pregunta']}»)")


# =====================================================================================
def escribir_preguntas():
    (PREGUNTAS / "preguntas.json").write_text(json.dumps(P, ensure_ascii=False, indent=1), encoding="utf-8")
    with open(PREGUNTAS / "preguntas.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["id", "fase", "workspace", "tipo", "dificultad", "rapida", "pregunta", "respuesta_esperada",
                    "fuente_esperada", "espera_rechazo", "revision_manual", "notas"])
        for x in P:
            w.writerow([x["id"], x["fase"], x["workspace"], x["tipo"], x["dificultad"], "sí" if x["rapida"] else "",
                        x["pregunta"], x["respuesta_esperada"],
                        " + ".join(x["fuente_esperada"]) if isinstance(x["fuente_esperada"], list) else x["fuente_esperada"],
                        "sí" if x["espera_rechazo"] else "", "sí" if x["revision_manual"] else "", x["notas"]])


def main():
    if "--solo-preguntas" in sys.argv:
        # no toca corpus/: reconstruirlo cambiaría las fechas internas de los .docx/.pptx/.pdf
        PREGUNTAS.mkdir(parents=True, exist_ok=True)
        escribir_preguntas()
        print(f"{len(P)} preguntas escritas en {PREGUNTAS} (corpus sin tocar)")
        return
    for d in [CORPUS, PREGUNTAS]:
        if d.exists() and d == CORPUS:
            shutil.rmtree(d)
        d.mkdir(parents=True, exist_ok=True)
    for d in [OFFICE, PDF_DIR, PDF_IMG, VAULT, USUARIO2, CAMBIOS]:
        d.mkdir(parents=True, exist_ok=True)

    docx_condor(OFFICE / f"{F_CON}.docx")
    pptx_sprint(OFFICE / f"{F_SPR}.pptx")
    pdf_concesion(PDF_DIR / f"{F_CRS}.pdf")
    pdf_procedimiento(PDF_DIR / f"{F_PRO}.pdf")
    pptx_fundicion(OFFICE / f"{F_FUN}.pptx")
    docx_clima(OFFICE / f"{F_CLI}.docx")
    pdf_escaneado_acta(PDF_IMG / f"{F_ACT}.pdf")
    docx_minuta(MI_RAG / f"{F_MIN}.docx")
    escribir_notas()

    (USUARIO2 / "Prueba_Aislamiento_Usuario2.md").write_text(
        "# Nota privada de Usuario2 (prueba de aislamiento)\n\n"
        "La palabra clave de aislamiento es GUAYACÁN-77. Si esta frase aparece en un espacio de trabajo de Usuario1, "
        "el aislamiento entre usuarios falló.\n\n" + AVISO + "\n", encoding="utf-8")

    # Fase F6: versiones modificadas (mismo nombre de archivo que el original)
    cam_office = CAMBIOS / "usuario1" / "Mi_RAG" / "office"
    cam_office.mkdir(parents=True, exist_ok=True)
    docx_condor(cam_office / f"{F_CON}.docx", version="2.4", fecha="2 de octubre de 2026", p95="300", jwt_min="30")
    cam_nota = CAMBIOS / "usuario1" / "Obsidian_Vault" / "Usuario1" / "Prueba RAG" / "Proyectos" / "Proyecto Cóndor.md"
    cam_nota.parent.mkdir(parents=True, exist_ok=True)
    original = NOTAS["Proyectos/Proyecto Cóndor.md"]
    linea = "- Próximo hito: salida a producción del módulo de certificados el **17 de noviembre de 2026**."
    assert linea in original
    cam_nota.write_text(original.replace(linea, NOTA_CONDOR_V2_HITO), encoding="utf-8")

    escribir_preguntas()

    n = len(P)
    print(f"Corpus generado en {CORPUS}")
    for f in sorted(CORPUS.rglob("*")):
        if f.is_file():
            print(f"  {f.relative_to(CORPUS)}  ({f.stat().st_size // 1024} KB)")
    print(f"Preguntas: {n} (rápidas: {sum(x['rapida'] for x in P)}) -> {PREGUNTAS}")


if __name__ == "__main__":
    main()
